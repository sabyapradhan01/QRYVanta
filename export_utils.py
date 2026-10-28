"""Safe, format-focused exports used by QRYVANTA's DOWNLOAD menu."""
from __future__ import annotations
import csv
from datetime import datetime
from pathlib import Path


def export_table_csv(path, headers, rows):
    """Write an exported Database Explorer result with a header row."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow([str(x) for x in headers])
        for row in rows:
            writer.writerow(["" if value is None else str(value) for value in row])
    return path


def format_command_history(history, app_name="QRYVANTA"):
    now = datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S %Z")
    lines = [f"{app_name} - SQL COMMAND HISTORY", f"Exported: {now}", f"Commands: {len(history)}", "=" * 72, ""]
    if not history:
        lines.append("No SQL statements have been captured in this session.")
    else:
        for i, command in enumerate(history, 1):
            lines += [f"-- Command {i:03d}", command.strip(), "", "-" * 72, ""]
    return "\n".join(lines).rstrip() + "\n"


def build_processing_report(report):
    """Render a report from the actual current app state; omit unavailable facts."""
    generated = datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S %Z")
    q = report.get("query", "") or "No SQL query has been analyzed yet."
    analysis = report.get("analysis") or {}
    findings = list(report.get("lint_findings") or [])
    perf = list(analysis.get("findings") or [])
    lines = [
        "QRYVANTA - QUERY PROCESSING REPORT",
        f"Generated: {generated}",
        f"Status: {report.get('status', 'Not available')}",
        f"Database type: {report.get('database_type') or 'Not available'}",
        "",
        "QUERY",
        "-----",
        q,
        "",
    ]
    if report.get("error"):
        lines += ["ANALYSIS ERROR", "--------------", str(report["error"]), ""]
    if analysis:
        lines += ["SUMMARY", "-------"]
        score = analysis.get("query_score")
        if score is not None:
            lines.append(f"Query health score: {score} / 100")
        query_info = analysis.get("query_info") or {}
        if isinstance(query_info, dict):
            for key in ("tables", "table_count", "join_count", "selected_column_count", "filter_count", "has_group_by", "has_order_by", "has_distinct", "aggregate_count", "where"):
                value = query_info.get(key)
                if value is not None:
                    lines.append(f"{key.replace('_', ' ').title()}: {value}")
        lines += ["", "STATIC LINT FINDINGS", "--------------------"]
        if not findings:
            lines.append("No static lint findings were returned.")
        for item in findings:
            lines.append(f"[{item.get('rule_id', 'RULE')}] {item.get('severity', 'Info')}: {item.get('message', '')}")
            if item.get("explanation"):
                lines.append(f"  Explanation: {item['explanation']}")
            if item.get("recommendation"):
                lines.append(f"  Recommendation: {item['recommendation']}")
        lines += ["", "PERFORMANCE FINDINGS", "--------------------"]
        if not perf:
            lines.append("No performance findings were returned by the available plan/context checks.")
        for item in perf:
            lines.append(f"[{item.get('rule_id', 'P')}] {item.get('severity', 'Info')}: {item.get('message', '')}")
            if item.get("explanation"):
                lines.append(f"  Explanation: {item['explanation']}")
            if item.get("recommendation"):
                lines.append(f"  Recommendation: {item['recommendation']}")
        lines += ["", "EXECUTION PLAN (AS REPORTED BY QRYVANTA)", "----------------------------------------"]
        lines.append(str(analysis.get("execution_plan") or "No execution plan text was available."))
        lines += ["", "INTERPRETATION NOTES", "--------------------",
                  "This report contains only data returned to QRYVANTA at export time. A score or finding is not a measured runtime. No execution duration or actual-row statistic is inferred unless the source data explicitly supplied it.",
                  "A missing finding is not proof that a query is optimal. Validate result correctness and compare the database's native plan and representative measurements before deploying a rewrite."]
    else:
        lines += ["No completed analysis data is available in the current session.",
                  "Run a supported SQL statement after connecting, then export the report again."]
    return "\n".join(lines) + "\n"


def save_report_txt(path, report):
    path = Path(path)
    path.write_text(build_processing_report(report), encoding="utf-8")
    return path


def save_report_pdf(path, report):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Preformatted, KeepTogether
    from xml.sax.saxutils import escape

    path = Path(path)
    text = build_processing_report(report)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportTitle", parent=styles["Title"], alignment=TA_CENTER, fontName="Helvetica-Bold", fontSize=16, leading=20, textColor=colors.HexColor("#14233C"), spaceAfter=8))
    styles.add(ParagraphStyle(name="ReportHeading", parent=styles["Heading2"], fontSize=11, leading=14, textColor=colors.HexColor("#214C78"), spaceBefore=10, spaceAfter=4))
    styles.add(ParagraphStyle(name="ReportBody", parent=styles["BodyText"], fontSize=8.5, leading=11, spaceAfter=3, wordWrap="CJK"))
    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=17*mm, leftMargin=17*mm, topMargin=16*mm, bottomMargin=16*mm, title="QRYVANTA Query Processing Report", author="QRYVANTA")
    story = []
    for line in text.splitlines():
        if line.startswith("QRYVANTA - "):
            story.append(Paragraph(escape(line), styles["ReportTitle"]))
        elif line and len(line) > 2 and set(line) == {"-"}:
            continue
        elif line.isupper() and line and len(line) <= 45:
            story.append(Spacer(1, 4))
            story.append(Paragraph(escape(line), styles["ReportHeading"]))
        elif not line:
            story.append(Spacer(1, 3))
        else:
            story.append(Paragraph(escape(line).replace("  ", "&nbsp;&nbsp;"), styles["ReportBody"]))
    doc.build(story)
    return path


def save_report_docx(path, report):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    path = Path(path)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)
    title = document.add_heading("QRYVANTA - Query Processing Report", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(f"Generated: {datetime.now().astimezone().strftime('%Y-%m-%d %H:%M:%S %Z')}")
    text = build_processing_report(report)
    for line in text.splitlines()[2:]:
        if line and len(line) <= 45 and line.isupper():
            document.add_heading(line.title(), level=1)
        elif line and set(line) == {"-"}:
            continue
        elif not line:
            continue
        else:
            para = document.add_paragraph(line)
            para.paragraph_format.space_after = Pt(3)
    footer = section.footer.paragraphs[0]
    footer.text = "QRYVANTA | Report reflects information available at export time."
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path
