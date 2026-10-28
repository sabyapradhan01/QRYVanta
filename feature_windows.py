"""Integrated offline learning, manual, quiz, exam practice and developer dialogs.

Optional media/photo assets live in assets/learn.mp4, assets/Dev1.jpg,
assets/Dev2.jpg and assets/fac.jpg. Missing files are handled gracefully.
"""
from __future__ import annotations
import html
import os
import random
import sys
from pathlib import Path

from PySide6.QtCore import Qt, QUrl
from PySide6.QtGui import QPixmap
from PySide6.QtWidgets import (
    QDialog, QHBoxLayout, QVBoxLayout, QGridLayout, QLabel, QListWidget,
    QListWidgetItem, QTextBrowser, QPushButton, QLineEdit, QFormLayout,
    QSpinBox, QComboBox, QRadioButton, QButtonGroup, QGroupBox,
    QMessageBox, QProgressBar, QFileDialog, QSplitter, QWidget, QStackedWidget,
    QFrame, QScrollArea, QTabWidget, QDialogButtonBox, QSizePolicy,
)

from learning_content import (
    LEARN_SECTIONS, SQL_RULES, PERFORMANCE_RULES, GLOSSARY, REFERENCES,
    QUESTION_BANK, COMPETITIVE_QUESTIONS,
)
from quiz_engine import build_attempt, validate_answer, summarize_attempt, available_questions, DIFFICULTIES

# Developer/faculty details displayed by the DEVELOPED BY dialog.
FACULTY_NAME = "Dr. Swaminathan A"
FACULTY_LABEL = "DBMS"
FACULTY_PHOTO = "assets/fac.jpg"
DEVELOPER_NAMES = ("Sabyasaachi Pradhan", "Subhashish Roy")
DEVELOPER_REGISTRATION_NUMBERS = ("25BCE1914", "25BCE5327")


def asset_path(relative: str) -> Path:
    """Resolve source and PyInstaller onedir/onefile asset paths."""
    base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
    candidates = [base / relative, Path(__file__).resolve().parent / relative]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return candidates[0]


def _dialog_style(theme):
    bg = theme.get("bg", "#F5F7FB")
    surface = theme.get("surface", "#FFFFFF")
    surface2 = theme.get("surface_2", "#EAF0F7")
    ink = theme.get("ink", "#152033")
    muted = theme.get("muted", "#526176")
    line = theme.get("line", "#CBD5E1")
    accent = theme.get("accent", "#375FEA")
    accent_ink = theme.get("accent_ink", "#FFFFFF")
    radius = int(theme.get("radius", 8))
    return f"""
        QDialog, QWidget#feature_root {{ background: {bg}; color: {ink}; }}
        QLabel {{ color: {ink}; }}
        QListWidget, QTextBrowser, QLineEdit, QSpinBox, QComboBox {{
            background: {surface}; color: {ink}; border: 1px solid {line};
            border-radius: {max(3, min(radius, 10))}px; padding: 7px;
            selection-background-color: {accent}; selection-color: {accent_ink};
        }}
        QListWidget::item {{ padding: 8px; }}
        QListWidget::item:selected {{ background: {accent}; color: {accent_ink}; }}
        QPushButton {{ background: {surface2}; color: {ink}; border: 1px solid {line};
            border-radius: {max(3, min(radius, 10))}px; padding: 8px 12px; font-weight: 700; }}
        QPushButton:hover {{ border-color: {accent}; }}
        QPushButton:disabled {{ color: {muted}; }}
        QGroupBox {{ color: {ink}; border: 1px solid {line}; border-radius: 6px; margin-top: 10px; padding-top: 8px; font-weight: 700; }}
        QGroupBox::title {{ subcontrol-origin: margin; left: 10px; padding: 0 4px; }}
        QRadioButton {{ color: {ink}; spacing: 7px; padding: 5px; }}
        QProgressBar {{ background: {surface2}; color: {ink}; border: 1px solid {line}; border-radius: 5px; text-align: center; }}
        QProgressBar::chunk {{ background: {accent}; border-radius: 4px; }}
    """


class FeatureDialog(QDialog):
    def __init__(self, title, theme, parent=None, width=1100, height=760):
        super().__init__(parent)
        self.theme = theme
        self.setWindowTitle(f"QRYVANTA | {title}")
        self.resize(width, height)
        self.setMinimumSize(min(width, 850), min(height, 600))
        self.setStyleSheet(_dialog_style(theme))

    def rich_text(self, content: str, title: str = "") -> str:
        ink = self.theme.get("ink", "#152033")
        muted = self.theme.get("muted", "#526176")
        accent = self.theme.get("accent", "#375FEA")
        return f"""<html><head><meta charset='utf-8'><style>
        body {{ color: {ink}; font-family: 'Segoe UI', Arial, sans-serif; font-size: 10.5pt; line-height: 1.45; }}
        h1 {{ color: {accent}; font-size: 19pt; margin-bottom: 5px; }}
        h2 {{ color: {accent}; font-size: 14pt; margin-top: 16px; }}
        h3 {{ font-size: 11.5pt; margin-top: 12px; }}
        p {{ margin: 5px 0 9px 0; }}
        pre, code {{ font-family: Consolas, 'Courier New', monospace; background: {self.theme.get('surface_2', '#EAF0F7')}; }}
        pre {{ padding: 9px; border: 1px solid {self.theme.get('line', '#CBD5E1')}; white-space: pre-wrap; }}
        li {{ margin: 3px 0; }} a {{ color: {accent}; }}
        .meta {{ color: {muted}; font-size: 9pt; }}
        .callout {{ border-left: 4px solid {accent}; padding: 7px 10px; background: {self.theme.get('surface_2', '#EAF0F7')}; }}
        </style></head><body>{('<h1>'+html.escape(title)+'</h1>') if title else ''}{content}</body></html>"""


class LearnDialog(FeatureDialog):
    """Navigation-based offline course, source-derived rule catalogue and glossary."""
    def __init__(self, theme, parent=None):
        super().__init__("LEARN", theme, parent, 1180, 800)
        root = QVBoxLayout(self)
        intro = QLabel("LEARN · SQL LINTING & QUERY PERFORMANCE")
        intro.setStyleSheet(f"font-size: 17px; font-weight: 900; color: {theme.get('accent', '#375FEA')};")
        root.addWidget(intro)
        splitter = QSplitter(Qt.Horizontal)
        self.navigation = QListWidget()
        self.page = QTextBrowser()
        self.page.setOpenExternalLinks(True)
        self.navigation.setMinimumWidth(250)
        self.navigation.setMaximumWidth(365)
        splitter.addWidget(self.navigation)
        splitter.addWidget(self.page)
        splitter.setStretchFactor(1, 1)
        root.addWidget(splitter, 1)
        footer = QHBoxLayout()
        self.progress_label = QLabel("Section 1")
        self.previous_button = QPushButton("← PREVIOUS")
        self.next_button = QPushButton("NEXT →")
        glossary_button = QPushButton("GLOSSARY")
        self.previous_button.clicked.connect(self.previous_page)
        self.next_button.clicked.connect(self.next_page)
        glossary_button.clicked.connect(self.show_glossary)
        footer.addWidget(self.progress_label)
        footer.addStretch()
        footer.addWidget(glossary_button)
        footer.addWidget(self.previous_button)
        footer.addWidget(self.next_button)
        root.addLayout(footer)

        self.pages = []
        for section in LEARN_SECTIONS:
            self.pages.append((section["title"], self.rich_text("<p>" + "</p><p>".join(html.escape(p).replace("\n", "<br>") for p in section["body"].split("\n\n")) + "</p>", section["title"])))
        for rule_id, title, severity, trigger, explanation, recommendation in SQL_RULES:
            body = self._rule_html("Static lint", rule_id, title, severity, trigger, explanation, recommendation)
            self.pages.append((f"{rule_id} · {title}", self.rich_text(body, f"{rule_id} · {title}")))
        for rule_id, title, severity, trigger, explanation, recommendation in PERFORMANCE_RULES:
            body = self._rule_html("Performance analysis", rule_id, title, severity, trigger, explanation, recommendation)
            self.pages.append((f"{rule_id} · {title}", self.rich_text(body, f"{rule_id} · {title}")))
        refs = "<h2>Credible references</h2><ul>" + "".join(
            f"<li><a href='{html.escape(url, quote=True)}'>{html.escape(name)}</a> — {html.escape(description)}</li>"
            for name, url, description in REFERENCES
        ) + "</ul><p>Use the documentation version matching your database engine and release. The rule catalogue documents the current source implementation, not every possible SQL or optimizer behavior.</p>"
        self.pages.append(("References", self.rich_text(refs, "References and further reading")))
        for title, _ in self.pages:
            self.navigation.addItem(QListWidgetItem(title))
        self.navigation.currentRowChanged.connect(self.show_page)
        self.navigation.setCurrentRow(0)
        self._add_video_panel(root)

    @staticmethod
    def _rule_html(kind, rule_id, title, severity, trigger, explanation, recommendation):
        return (f"<p><b>Category:</b> {html.escape(kind)} &nbsp; <b>Severity in source:</b> {html.escape(severity)}</p>"
                f"<h2>What the code checks</h2><p>{html.escape(trigger)}</p>"
                f"<h2>Why it matters</h2><p>{html.escape(explanation)}</p>"
                f"<h2>Recommended response</h2><p>{html.escape(recommendation)}</p>"
                "<p class='callout'><b>Interpretation:</b> this is a rule signal. Check the query's intent, database dialect, plan and workload before changing SQL or indexes.</p>")

    def show_page(self, row):
        if row < 0 or row >= len(self.pages):
            return
        self.page.setHtml(self.pages[row][1])
        self.progress_label.setText(f"Section {row + 1} of {len(self.pages)}")
        self.previous_button.setEnabled(row > 0)
        self.next_button.setEnabled(row < len(self.pages) - 1)

    def previous_page(self):
        self.navigation.setCurrentRow(max(0, self.navigation.currentRow() - 1))

    def next_page(self):
        self.navigation.setCurrentRow(min(len(self.pages) - 1, self.navigation.currentRow() + 1))

    def show_glossary(self):
        body = "<dl>" + "".join(f"<dt><b>{html.escape(k)}</b></dt><dd>{html.escape(v)}</dd>" for k, v in sorted(GLOSSARY.items())) + "</dl>"
        dialog = FeatureDialog("GLOSSARY", self.theme, self, 720, 640)
        layout = QVBoxLayout(dialog)
        text = QTextBrowser()
        text.setHtml(dialog.rich_text(body, "SQL and Performance Glossary"))
        layout.addWidget(text)
        close = QPushButton("CLOSE")
        close.clicked.connect(dialog.accept)
        layout.addWidget(close, 0, Qt.AlignRight)
        dialog.exec()

    def _add_video_panel(self, root):
        group = QGroupBox("LOCAL VIDEO LESSON")
        layout = QHBoxLayout(group)
        self.video_status = QLabel()
        self.video_status.setWordWrap(True)
        self.video_status.setMinimumWidth(280)
        self.video_status.setText("Optional lesson video: assets/learn.mp4")
        layout.addWidget(self.video_status, 1)
        self.video_widget = None
        self.media_player = None
        self.play_button = QPushButton("PLAY VIDEO")
        self.play_button.setEnabled(False)
        layout.addWidget(self.play_button)
        try:
            from PySide6.QtMultimedia import QMediaPlayer, QAudioOutput
            from PySide6.QtMultimediaWidgets import QVideoWidget
            video_path = asset_path("assets/learn.mp4")
            if not video_path.is_file():
                self.video_status.setText("Video not included in this source ZIP. Add your own licensed video at assets/learn.mp4. All text lessons work offline without it.")
            else:
                self.video_widget = QVideoWidget()
                self.video_widget.setMinimumHeight(120)
                self.video_widget.setMaximumHeight(180)
                layout.insertWidget(0, self.video_widget, 2)
                self.media_player = QMediaPlayer(self)
                audio = QAudioOutput(self)
                self.media_player.setAudioOutput(audio)
                self.media_player.setVideoOutput(self.video_widget)
                self.media_player.errorOccurred.connect(self.handle_media_error)
                self.media_player.setSource(QUrl.fromLocalFile(str(video_path)))
                self.play_button.setEnabled(True)
                self.play_button.setText("PLAY / PAUSE")
                self.play_button.clicked.connect(self.toggle_video)
                self.video_status.setText(f"Loaded local video: {video_path.name}. It is played from disk and does not require a network connection.")
        except Exception as error:
            self.video_status.setText(f"Local video playback is unavailable in this Qt build ({type(error).__name__}). The complete text course remains available offline. Place optional video at assets/learn.mp4 when Qt Multimedia is installed.")
        root.addWidget(group)

    def handle_media_error(self, _error, error_string):
        self.video_status.setText(
            "The optional local video could not be played. "
            f"{error_string or 'Check the file and the installed media codecs.'} "
            "The complete text course remains available offline."
        )
        self.play_button.setEnabled(False)

    def toggle_video(self):
        if self.media_player is None:
            return
        from PySide6.QtMultimedia import QMediaPlayer
        if self.media_player.playbackState() == QMediaPlayer.PlayingState:
            self.media_player.pause()
        else:
            self.media_player.play()


HELP_SECTIONS = [
    ("Overview", "QRYVANTA combines SQL execution, static linting, performance-plan interpretation, Database Explorer, account management, five themes, study content and exports. Database features need a configured database/driver; learning, manual, quiz and exam practice use local packaged content."),
    ("ANALYZER", "Select ANALYZER to show the SQL Console beside the analysis panel. When a supported SQL statement is completed, QRYVANTA parses it with SQLGlot for the active database dialect, runs static rules, builds available query/plan context and displays score, findings and recommendations. Analysis depends on a successful login and database adapter. The health score summarizes finding severity; it is not elapsed time. The lint and performance areas list findings, while recommendation cards are suggestions that should be checked against the real plan."),
    ("Status labels", "The connection status in the top bar identifies whether the app is connected/login-required and shows the active database/account label after connection. The bottom status label reports major application states such as ready, waiting for login, analyzing, analysis complete or analysis failed. Detailed parse/database errors appear in the analysis panel or SQL console; do not assume a status label means a query completed successfully."),
    ("DATABASE", "Select DATABASE to show Database Explorer. The left list contains tables available through the connected account. Click a table to load its rows and columns; use REFRESH to ask the adapter for current table metadata again. Empty results can mean the table has no rows or the account lacks visibility. CSV export includes the headers and the currently displayed result, not an unrequested full database dump."),
    ("SQL Console", "Type or paste SQL in the lower editor and submit the statement as supported by the console. Multi-line input is accumulated until a semicolon outside quoted strings is reached; the capture helper treats SQL statements beginning with common SELECT/WITH/INSERT/UPDATE/DELETE/MERGE/DDL keywords as analyzable statements. Oracle console commands are sent to SQL*Plus; supported non-Oracle adapters execute through their adapter. The output area is read-only."),
    ("SQL command history", "Use the input's Up/Down history navigation to revisit statements captured in the current running session. DOWNLOAD → SQL Command History (TXT) saves the captured SQL statements with numbering and separators. History is in-memory; it is not a promise of long-term archival. Exported SQL may contain schema or business details, so store the file carefully."),
    ("LOGIN / SWITCH ACCOUNT / LOGOUT", "LOGIN opens database profile setup and connection testing. Provide the database type, host, port, database/service, username and password from an authorized administrator. SWITCH ACCOUNT opens saved profiles and supports adding/removing/switching profiles. LOGOUT disconnects and clears the active session before the application returns to login. Save-password support uses the OS credential store; never distribute credential stores or secrets."),
    ("S|S and S↕S", "S|S switches the central splitter to side-by-side (console at left, current right-side view at right). S↕S switches to a vertical stacked layout. The button representing the active layout may be hidden; the other button is the action to switch to."),
    ("THEME and Day/Night", "THEME selects Neo-Brutalism, Cybercore, Swiss, Bento Grid or Cyberpunk. Day/Night switches the shared light/dark palette while retaining the chosen theme's typography and accent identity; it is persisted with application settings. If a custom widget has hard-coded styling or a third-party control looks inconsistent, report its screen and theme so that selector can be corrected."),
    (("? menu"), "The question-mark menu opens LEARN (offline SQL course plus source-derived SQL001–SQL015 and P001–P022 catalogue, glossary and optional local video), HELP (this complete manual), QUIZ (randomized offline quiz), COMPETITIVE EXAM (original solved SQL practice, not mislabeled past papers), and DEVELOPED BY (team/faculty information)."),
    ("DOWNLOAD menu", "DOWNLOAD → Database Table CSV asks for a file and exports the currently visible Database Explorer table, including the column headers. SQL Command History (TXT) writes captured statements in a numbered format. Processing Report supports PDF, DOCX and TXT and includes the latest analysis information available in memory; it does not fabricate a duration or actual row count. If no completed analysis exists, the report says so."),
    ("QUIZ", "Choose a question count and Easy/Medium/Hard. The UI updates the allowed count to the selected difficulty's available bank size. Starting an attempt randomly samples unique question IDs and shuffles each question's options. Select an option and choose CHECK ANSWER; each item is evaluated once, displays the required feedback sentence and an explanation. NEXT moves through the attempt. The final page reports score, percentage, attempted/unanswered count and topic-level performance."),
    ("COMPETITIVE EXAM", "This is an offline bank of 50+ original competitive-exam-style practice questions with answers, explanations and topic tags. Questions are explicitly labeled as original practice material, not actual past-paper questions. No exam year/authority is asserted without a verified citation. Search/filter the bank and use the export buttons for questions-only or solutions documents."),
    ("LEARN and local assets", "LEARN's text lessons and rule descriptions are packaged offline. `assets/learn.mp4` is optional: if absent, playback is replaced by a clear instruction; if Qt Multimedia/codecs are unavailable, text lessons remain usable. `assets/Dev1.jpg`, `assets/Dev2.jpg`, and `assets/fac.jpg` are optional; missing pictures use initials/placeholders. Add approved images to those exact paths before rebuilding."),
    ("Processing reports", "The report stores status, current query, active database type, query score if returned, static and performance findings, and available execution-plan text. Analysis errors are included transparently. It deliberately says when fields are unavailable. A report score is not the measured runtime; use native plan tools and repeatable benchmarks for tuning decisions."),
    ("Troubleshooting: login and adapters", "Connection refused/timeouts: confirm server, host, port, firewall/VPN and service/database. Authentication errors: confirm account/password with the administrator. Permission errors: ask for minimal required privileges. Confirm the correct Oracle service name, not an assumed SID. MySQL/PostgreSQL need their compatible Python connector packages. Avoid embedding passwords in source or README."),
    ("Troubleshooting: SQL*Plus", "Oracle's SQL console needs a compatible SQL*Plus installation. If prompted, choose the actual sqlplus.exe file. Confirm the SQL*Plus client matches installed Oracle libraries and can connect independently. A database adapter connection can succeed even when the external SQL*Plus console is missing or misconfigured."),
    ("Troubleshooting: analysis and exports", "A parse error may indicate unsupported dialect syntax or a statement that is not one complete SQL statement. Test against the active database dialect. If no findings appear, it means configured rules did not match available syntax/plan evidence—not that the query is optimal. CSV exports only displayed rows. Install the optional `reportlab` and `python-docx` dependencies for PDF/DOCX report export; TXT and CSV use the Python standard library."),
    ("Security and safe operation", "Use only databases you are authorized to access. Confirm the active account/database before data-changing SQL. Use least privilege, parameterized statements in apps, and backups/transactions appropriate to the data. Before sharing logs, SQL history or reports, remove sensitive schema values and business data. Do not include `.env`, local database credentials or credential-store data in source packages."),
    ("Configuration and files", "Main UI/orchestration: app.py. Database Explorer: database_view.py. Connectors: adapters/*.py and core/database.py. Static lint rules: linter/rules.py and linter/linter.py. Performance rules: analyzer/performance_rules.py and analyzer/analyzer.py. Learning, rule descriptions and canonical quiz bank: learning_content.py. Quiz state/selection/feedback: quiz_engine.py. Feature dialogs: feature_windows.py. Export functions: export_utils.py. PyInstaller configuration: QRYVANTA.spec; Inno Setup installer: QRYVANTA-Setup.iss. Faculty defaults are documented in feature_windows.py."),
]


class HelpDialog(FeatureDialog):
    def __init__(self, theme, parent=None):
        super().__init__("HELP", theme, parent, 1120, 780)
        layout = QVBoxLayout(self)
        heading = QLabel("HELP MANUAL · CONTROLS, OPERATIONS & TROUBLESHOOTING")
        heading.setStyleSheet(f"font-size: 16px; font-weight: 900; color: {theme.get('accent', '#375FEA')};")
        layout.addWidget(heading)
        splitter = QSplitter(Qt.Horizontal)
        self.navigation = QListWidget()
        self.text = QTextBrowser()
        self.text.setOpenExternalLinks(True)
        splitter.addWidget(self.navigation)
        splitter.addWidget(self.text)
        splitter.setStretchFactor(1, 1)
        layout.addWidget(splitter, 1)
        for title, _ in HELP_SECTIONS:
            self.navigation.addItem(title)
        self.navigation.currentRowChanged.connect(self.show_section)
        self.navigation.setCurrentRow(0)
        bottom = QHBoxLayout()
        bottom.addStretch()
        close = QPushButton("CLOSE")
        close.clicked.connect(self.accept)
        bottom.addWidget(close)
        layout.addLayout(bottom)

    def show_section(self, row):
        if not 0 <= row < len(HELP_SECTIONS):
            return
        title, body = HELP_SECTIONS[row]
        paragraphs = "".join(f"<p>{html.escape(part)}</p>" for part in body.split("\n\n"))
        self.text.setHtml(self.rich_text(paragraphs, title))


class QuizDialog(FeatureDialog):
    def __init__(self, theme, parent=None):
        super().__init__("QUIZ", theme, parent, 940, 730)
        self.bank = QUESTION_BANK
        self.questions = []
        self.responses = {}
        self.cursor = 0
        self.checked = set()
        self.current_radio_group = QButtonGroup(self)
        root = QVBoxLayout(self)
        self.stack = QStackedWidget()
        root.addWidget(self.stack, 1)
        self.setup_screen = QWidget()
        setup_layout = QVBoxLayout(self.setup_screen)
        title = QLabel("OFFLINE SQL KNOWLEDGE QUIZ")
        title.setStyleSheet(f"font-size: 20px; font-weight: 900; color: {theme.get('accent', '#375FEA')};")
        setup_layout.addWidget(title)
        setup_layout.addWidget(QLabel("Questions are sampled without replacement. Question and option order are randomized each attempt."))
        form = QFormLayout()
        self.difficulty_combo = QComboBox()
        self.difficulty_combo.addItems(list(DIFFICULTIES))
        self.count_spin = QSpinBox()
        self.count_spin.setRange(1, len(available_questions(self.bank, "Easy")))
        self.count_spin.setValue(min(10, self.count_spin.maximum()))
        form.addRow("Difficulty", self.difficulty_combo)
        form.addRow("Question count", self.count_spin)
        setup_layout.addLayout(form)
        self.available_label = QLabel()
        setup_layout.addWidget(self.available_label)
        self.difficulty_combo.currentTextChanged.connect(self.update_count_limit)
        self.update_count_limit()
        self.start_button = QPushButton("START RANDOMIZED ATTEMPT")
        self.start_button.clicked.connect(self.start_attempt)
        setup_layout.addWidget(self.start_button)
        setup_layout.addStretch()
        self.stack.addWidget(self.setup_screen)

        self.question_screen = QWidget()
        q_layout = QVBoxLayout(self.question_screen)
        self.progress = QProgressBar()
        q_layout.addWidget(self.progress)
        self.question_heading = QLabel()
        self.question_heading.setWordWrap(True)
        self.question_heading.setStyleSheet("font-size: 17px; font-weight: 800;")
        q_layout.addWidget(self.question_heading)
        self.meta_label = QLabel()
        q_layout.addWidget(self.meta_label)
        self.options_group = QGroupBox("Choose one answer")
        options_layout = QVBoxLayout(self.options_group)
        self.option_radios = []
        for _ in range(4):
            radio = QRadioButton()
            self.current_radio_group.addButton(radio)
            options_layout.addWidget(radio)
            self.option_radios.append(radio)
        q_layout.addWidget(self.options_group)
        self.feedback_label = QLabel("")
        self.feedback_label.setWordWrap(True)
        self.feedback_label.setMinimumHeight(70)
        q_layout.addWidget(self.feedback_label)
        buttons = QHBoxLayout()
        self.check_button = QPushButton("CHECK ANSWER")
        self.next_button = QPushButton("NEXT QUESTION →")
        self.finish_button = QPushButton("FINISH ATTEMPT")
        self.check_button.clicked.connect(self.check_current_answer)
        self.next_button.clicked.connect(self.next_question)
        self.finish_button.clicked.connect(self.finish_attempt)
        buttons.addWidget(self.check_button)
        buttons.addStretch()
        buttons.addWidget(self.next_button)
        buttons.addWidget(self.finish_button)
        q_layout.addLayout(buttons)
        self.stack.addWidget(self.question_screen)

        self.summary_screen = QWidget()
        summary_layout = QVBoxLayout(self.summary_screen)
        self.summary_heading = QLabel("QUIZ SUMMARY")
        self.summary_heading.setStyleSheet(f"font-size: 22px; font-weight: 900; color: {theme.get('accent', '#375FEA')};")
        summary_layout.addWidget(self.summary_heading)
        self.summary_text = QTextBrowser()
        summary_layout.addWidget(self.summary_text, 1)
        again = QPushButton("NEW ATTEMPT")
        again.clicked.connect(lambda: self.stack.setCurrentWidget(self.setup_screen))
        summary_layout.addWidget(again)
        self.stack.addWidget(self.summary_screen)

    def update_count_limit(self, *_args):
        difficulty = self.difficulty_combo.currentText()
        total = len(available_questions(self.bank, difficulty))
        value = min(self.count_spin.value(), total)
        self.count_spin.blockSignals(True)
        self.count_spin.setRange(1, max(1, total))
        self.count_spin.setValue(max(1, value))
        self.count_spin.blockSignals(False)
        self.available_label.setText(f"{total} questions currently available at {difficulty} difficulty.")

    def start_attempt(self):
        try:
            self.questions = build_attempt(self.bank, self.count_spin.value(), self.difficulty_combo.currentText())
        except Exception as error:
            QMessageBox.warning(self, "Cannot start quiz", str(error))
            return
        self.responses = {}
        self.checked = set()
        self.cursor = 0
        self.show_question()
        self.stack.setCurrentWidget(self.question_screen)

    def show_question(self):
        question = self.questions[self.cursor]
        self.progress.setRange(0, len(self.questions))
        self.progress.setValue(self.cursor + 1)
        self.question_heading.setText(f"Question {self.cursor + 1} of {len(self.questions)}\n\n{question['question']}")
        self.meta_label.setText(f"{question['difficulty']} · {question['topic']} · ID {question['id']}")
        self.current_radio_group.setExclusive(False)
        for radio in self.option_radios:
            radio.setChecked(False)
            radio.setEnabled(question['id'] not in self.checked)
        self.current_radio_group.setExclusive(True)
        for idx, radio in enumerate(self.option_radios):
            if idx < len(question['options']):
                radio.setText(question['options'][idx])
                radio.show()
            else:
                radio.hide()
        answered = question['id'] in self.checked
        self.feedback_label.setText("")
        self.feedback_label.setStyleSheet("")
        self.check_button.setEnabled(not answered)
        self.next_button.setEnabled(answered)
        if self.cursor == len(self.questions) - 1:
            self.next_button.setText("VIEW FINAL SCORE →")
        else:
            self.next_button.setText("NEXT QUESTION →")
        self.finish_button.setEnabled(True)
        if answered:
            selected = self.responses.get(question['id'])
            is_correct, message, explanation = validate_answer(question, selected)
            self.feedback_label.setText(f"{message}\nCorrect answer: {question['answer']}\n{explanation}")

    def check_current_answer(self):
        question = self.questions[self.cursor]
        if question['id'] in self.checked:
            return
        selected = None
        for radio in self.option_radios:
            if radio.isChecked():
                selected = radio.text()
                break
        self.responses[question['id']] = selected
        self.checked.add(question['id'])
        is_correct, message, explanation = validate_answer(question, selected)
        self.feedback_label.setText(f"{message}\nCorrect answer: {question['answer']}\nExplanation: {explanation}")
        self.feedback_label.setStyleSheet(f"font-weight: 700; color: {self.theme.get('ok_fg', '#147D49') if is_correct else self.theme.get('bad_fg', '#B4233C')};")
        for radio in self.option_radios:
            radio.setEnabled(False)
        self.check_button.setEnabled(False)
        self.next_button.setEnabled(True)
        if self.cursor == len(self.questions) - 1:
            self.next_button.setText("VIEW FINAL SCORE →")

    def next_question(self):
        if self.cursor >= len(self.questions) - 1:
            self.finish_attempt()
            return
        self.cursor += 1
        self.next_button.setText("NEXT QUESTION →")
        self.show_question()

    def finish_attempt(self):
        if not self.questions:
            self.stack.setCurrentWidget(self.setup_screen)
            return
        summary = summarize_attempt(self.questions, self.responses)
        by_topic = "".join(f"<li><b>{html.escape(topic)}</b>: {data['correct']}/{data['total']} correct</li>" for topic, data in sorted(summary['by_topic'].items()))
        body = (f"<h2>{summary['correct']} / {summary['total']} correct ({summary['percentage']:.1f}%)</h2>"
                f"<p>Attempted: {summary['attempted']} · Incorrect or unanswered: {summary['incorrect_or_unanswered']}</p>"
                f"<h3>Topic breakdown</h3><ul>{by_topic}</ul>"
                "<p>Use explanations to review concepts. These are original practice questions and are not labeled as actual past-paper items.</p>")
        temp = FeatureDialog("SCORE", self.theme, self)
        self.summary_text.setHtml(temp.rich_text(body, "Your performance"))
        self.stack.setCurrentWidget(self.summary_screen)


class CompetitiveExamDialog(FeatureDialog):
    def __init__(self, theme, parent=None):
        super().__init__("COMPETITIVE EXAM", theme, parent, 1180, 800)
        self.bank = COMPETITIVE_QUESTIONS
        self.filtered = list(self.bank)
        root = QVBoxLayout(self)
        heading = QLabel(f"COMPETITIVE EXAM PRACTICE · {len(self.bank)} SOLVED QUESTIONS")
        heading.setStyleSheet(f"font-size: 17px; font-weight: 900; color: {theme.get('accent', '#375FEA')};")
        root.addWidget(heading)
        note = QLabel("Original SQL practice questions, not attributed to actual past papers. No exam/year label is asserted without a verified source.")
        note.setWordWrap(True)
        root.addWidget(note)
        controls = QHBoxLayout()
        self.search = QLineEdit()
        self.search.setPlaceholderText("Search question text or topic...")
        self.difficulty = QComboBox()
        self.difficulty.addItems(["All difficulties", *DIFFICULTIES])
        controls.addWidget(self.search, 1)
        controls.addWidget(self.difficulty)
        root.addLayout(controls)
        splitter = QSplitter(Qt.Horizontal)
        self.question_list = QListWidget()
        self.detail = QTextBrowser()
        splitter.addWidget(self.question_list)
        splitter.addWidget(self.detail)
        splitter.setStretchFactor(1, 1)
        root.addWidget(splitter, 1)
        buttons = QHBoxLayout()
        export_q = QPushButton("EXPORT QUESTIONS DOCX")
        export_s = QPushButton("EXPORT SOLUTIONS DOCX")
        close = QPushButton("CLOSE")
        export_q.clicked.connect(lambda: self.export_docx(solutions=False))
        export_s.clicked.connect(lambda: self.export_docx(solutions=True))
        close.clicked.connect(self.accept)
        buttons.addWidget(export_q)
        buttons.addWidget(export_s)
        buttons.addStretch()
        buttons.addWidget(close)
        root.addLayout(buttons)
        self.search.textChanged.connect(self.apply_filter)
        self.difficulty.currentTextChanged.connect(self.apply_filter)
        self.question_list.currentRowChanged.connect(self.show_question)
        self.apply_filter()

    def apply_filter(self, *_args):
        query = self.search.text().strip().lower()
        difficulty = self.difficulty.currentText()
        self.filtered = [q for q in self.bank if (difficulty == "All difficulties" or q['difficulty'] == difficulty) and (not query or query in q['question'].lower() or query in q['topic'].lower() or query in q['id'].lower())]
        self.question_list.clear()
        for q in self.filtered:
            self.question_list.addItem(f"{q['id']} · {q['difficulty']} · {q['topic']} — {q['question']}")
        if self.filtered:
            self.question_list.setCurrentRow(0)
        else:
            self.detail.setPlainText("No questions match the current filter.")

    def show_question(self, row):
        if not 0 <= row < len(self.filtered):
            return
        q = self.filtered[row]
        options = "".join(f"<li>{html.escape(option)}</li>" for option in q['options'])
        body = (f"<p><b>Topic tags:</b> {html.escape(q['topic'])}, {html.escape(q['difficulty'])}</p>"
                f"<p>{html.escape(q['question'])}</p><ol>{options}</ol>"
                f"<h2>Answer</h2><p><b>{html.escape(q['answer'])}</b></p>"
                f"<h2>Explanation</h2><p>{html.escape(q['explanation'])}</p>"
                f"<p class='meta'>{html.escape(q.get('source_note', 'Original practice question.'))}</p>")
        self.detail.setHtml(self.rich_text(body, q['id']))

    def export_docx(self, solutions=False):
        default = "QRYVANTA_Competitive_SQL_Solutions.docx" if solutions else "QRYVANTA_Competitive_SQL_Questions.docx"
        path, _ = QFileDialog.getSaveFileName(self, "Export practice document", default, "Word document (*.docx)")
        if not path:
            return
        try:
            from docx import Document
            from docx.shared import Pt, Inches
            document = Document()
            document.sections[0].top_margin = Inches(0.65)
            document.sections[0].bottom_margin = Inches(0.65)
            document.styles['Normal'].font.name = 'Aptos'
            document.styles['Normal'].font.size = Pt(9)
            document.add_heading("QRYVANTA | Competitive SQL Practice", 0)
            document.add_paragraph(f"{len(self.bank)} original practice questions. Not represented as actual past-paper questions; exam/year attribution is omitted unless independently verified.")
            for i, q in enumerate(self.bank, 1):
                document.add_heading(f"{i}. [{q['id']}] {q['topic']} ({q['difficulty']})", level=2)
                document.add_paragraph(q['question'])
                for letter, option in zip("ABCD", q['options']):
                    document.add_paragraph(f"{letter}. {option}", style='List Bullet')
                if solutions:
                    para = document.add_paragraph()
                    para.add_run("Correct answer: ").bold = True
                    para.add_run(q['answer'])
                    para = document.add_paragraph()
                    para.add_run("Explanation: ").bold = True
                    para.add_run(q['explanation'])
                else:
                    document.add_paragraph("Answer: ______________________________________________")
                document.add_paragraph(f"Topic tags: {q['topic']} | {q['difficulty']} | Original practice material")
            document.save(path)
            QMessageBox.information(self, "Export complete", f"Saved:\n{path}")
        except Exception as error:
            QMessageBox.critical(self, "Export failed", f"Could not create DOCX. Install python-docx and retry.\n\n{error}")


class DeveloperDialog(FeatureDialog):
    """Team and faculty information with consistently aligned portrait cards."""

    def __init__(self, theme, parent=None):
        super().__init__("DEVELOPED BY", theme, parent, 1060, 700)
        root = QVBoxLayout(self)
        root.setContentsMargins(18, 14, 18, 14)
        root.setSpacing(12)

        accent = theme.get('accent', '#375FEA')
        line = theme.get('line', '#CBD5E1')
        surface2 = theme.get('surface_2', '#EAF0F7')
        muted = theme.get('muted', '#526176')

        title = QLabel("DEVELOPED BY")
        title.setStyleSheet(f"font-size: 22px; font-weight: 900; color: {accent};")
        root.addWidget(title)

        # Equal-height cards keep the image, name and registration number aligned,
        # whether a portrait exists or the fallback initials are being displayed.
        developers = QHBoxLayout()
        developers.setSpacing(14)
        for index, (name, registration_number) in enumerate(
            zip(DEVELOPER_NAMES, DEVELOPER_REGISTRATION_NUMBERS), 1
        ):
            card = QFrame()
            card.setObjectName("developerCard")
            card.setFrameShape(QFrame.StyledPanel)
            card.setMinimumHeight(255)
            card.setStyleSheet(
                f"QFrame#developerCard {{ border: 1px solid {line}; "
                f"border-radius: 8px; background: {theme.get('surface', '#FFFFFF')}; }}"
            )
            card_layout = QVBoxLayout(card)
            card_layout.setContentsMargins(14, 12, 14, 10)
            card_layout.setSpacing(6)

            image = QLabel()
            image.setFixedSize(136, 136)
            image.setAlignment(Qt.AlignCenter)
            image.setStyleSheet(
                f"border: 1px solid {line}; border-radius: 10px; "
                f"background: {surface2};"
            )
            photo = asset_path(f"assets/Dev{index}.jpg")
            pixmap = QPixmap(str(photo)) if photo.is_file() else QPixmap()
            has_photo = not pixmap.isNull()
            if has_photo:
                image.setPixmap(
                    pixmap.scaled(136, 136, Qt.KeepAspectRatio, Qt.SmoothTransformation)
                )
            else:
                image.setText(name.split()[0][:1].upper())
                image.setStyleSheet(
                    f"font-size: 48px; font-weight: 900; color: {theme.get('ink', '#152033')}; "
                    f"border: 1px solid {line}; border-radius: 10px; background: {surface2};"
                )

            name_label = QLabel(name)
            name_label.setAlignment(Qt.AlignCenter)
            name_label.setMinimumHeight(23)
            name_label.setStyleSheet(
                "font-size: 15px; font-weight: 800; border: none; background: transparent;"
            )

            registration_label = QLabel(f"Registration No.: {registration_number}")
            registration_label.setAlignment(Qt.AlignCenter)
            registration_label.setMinimumHeight(19)
            registration_label.setStyleSheet(
                f"font-size: 12px; color: {muted}; border: none; background: transparent;"
            )

            # Reserve one line on both cards so missing-photo text does not shift
            # the name or registration number vertically.
            photo_status = QLabel("" if has_photo else f"Photo not added · assets/Dev{index}.jpg")
            photo_status.setAlignment(Qt.AlignCenter)
            photo_status.setFixedHeight(17)
            photo_status.setStyleSheet(
                f"font-size: 10px; color: {muted}; border: none; background: transparent;"
            )

            card_layout.addWidget(image, 0, Qt.AlignHCenter)
            card_layout.addWidget(name_label)
            card_layout.addWidget(registration_label)
            card_layout.addWidget(photo_status)
            card_layout.addStretch(1)
            developers.addWidget(card, 1)
        root.addLayout(developers)

        # Faculty portrait and its two-line text are one centered horizontal unit.
        # Keep the group compact instead of letting the portrait and labels drift
        # to different vertical positions in a tall, mostly empty panel.
        faculty = QGroupBox("FACULTY")
        faculty.setMinimumHeight(205)
        faculty.setMaximumHeight(220)
        faculty_layout = QHBoxLayout(faculty)
        faculty_layout.setContentsMargins(22, 18, 22, 18)
        faculty_layout.setSpacing(18)
        faculty_layout.addStretch(1)

        faculty_image = QLabel()
        faculty_image.setFixedSize(144, 144)
        faculty_image.setAlignment(Qt.AlignCenter)
        faculty_image.setStyleSheet(
            f"border: 1px solid {line}; border-radius: 10px; background: {surface2};"
        )
        faculty_photo = asset_path(FACULTY_PHOTO)
        faculty_pixmap = QPixmap(str(faculty_photo)) if faculty_photo.is_file() else QPixmap()
        if not faculty_pixmap.isNull():
            faculty_image.setPixmap(
                faculty_pixmap.scaled(142, 142, Qt.KeepAspectRatio, Qt.SmoothTransformation)
            )
        else:
            faculty_image.setText("F")
            faculty_image.setStyleSheet(
                f"font-size: 48px; font-weight: 900; color: {theme.get('ink', '#152033')}; "
                f"border: 1px solid {line}; border-radius: 10px; background: {surface2};"
            )
        faculty_layout.addWidget(faculty_image, 0, Qt.AlignVCenter)

        faculty_details = QVBoxLayout()
        faculty_details.setContentsMargins(0, 0, 0, 0)
        faculty_details.setSpacing(8)
        faculty_details.setAlignment(Qt.AlignVCenter | Qt.AlignLeft)

        faculty_name_label = QLabel(FACULTY_NAME)
        faculty_name_label.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        faculty_name_label.setStyleSheet(
            "font-size: 19px; font-weight: 900; border: none; background: transparent;"
        )

        faculty_department_label = QLabel(FACULTY_LABEL)
        faculty_department_label.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        faculty_department_label.setStyleSheet(
            f"font-size: 15px; font-weight: 700; color: {accent}; "
            "border: none; background: transparent;"
        )

        faculty_details.addStretch(1)
        faculty_details.addWidget(faculty_name_label)
        faculty_details.addWidget(faculty_department_label)
        if not faculty_photo.is_file() or faculty_pixmap.isNull():
            missing_faculty_photo = QLabel(f"Photo not added · {FACULTY_PHOTO}")
            missing_faculty_photo.setWordWrap(True)
            missing_faculty_photo.setStyleSheet(
                f"font-size: 11px; color: {muted}; border: none; background: transparent;"
            )
            faculty_details.addWidget(missing_faculty_photo)
        faculty_details.addStretch(1)
        faculty_layout.addLayout(faculty_details, 0)
        faculty_layout.addStretch(1)
        root.addWidget(faculty)

        controls = QHBoxLayout()
        close = QPushButton("CLOSE")
        close.clicked.connect(self.accept)
        controls.addStretch()
        controls.addWidget(close)
        root.addLayout(controls)
