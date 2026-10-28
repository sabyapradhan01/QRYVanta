#!/usr/bin/env python3
"""Generate QRYVANTA offline study documents from learning_content.py.
Run from the project root: python tools/generate_study_documents.py
"""
from __future__ import annotations
import sys
from pathlib import Path
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from learning_content import LEARN_SECTIONS, SQL_RULES, PERFORMANCE_RULES, GLOSSARY, REFERENCES, QUESTION_BANK

OUTPUT = ROOT / "docs"
OUTPUT.mkdir(exist_ok=True)


def docx_base(title, subtitle):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    document = Document()
    sec = document.sections[0]
    sec.top_margin = Inches(0.68)
    sec.bottom_margin = Inches(0.68)
    sec.left_margin = Inches(0.76)
    sec.right_margin = Inches(0.76)
    styles = document.styles
    styles['Normal'].font.name = 'Aptos'
    styles['Normal'].font.size = Pt(9.5)
    styles['Normal'].paragraph_format.space_after = Pt(5)
    for style_name, size in [('Title', 28), ('Heading 1', 17), ('Heading 2', 12), ('Heading 3', 10)]:
        st = styles[style_name]
        st.font.name = 'Aptos Display' if style_name in ('Title', 'Heading 1') else 'Aptos'
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(32, 75, 120)
    title_p = document.add_paragraph(style='Title')
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.add_run(title)
    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(subtitle)
    run.italic = True
    run.font.color.rgb = RGBColor(88, 101, 120)
    document.add_paragraph('')
    footer = sec.footer.paragraphs[0]
    footer.text = 'QRYVANTA | Query Intelligence Studio | Original learning material'
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return document


def add_paragraphs(doc, body):
    for paragraph in body.split('\n\n'):
        p = doc.add_paragraph()
        for i, line in enumerate(paragraph.split('\n')):
            if i:
                p.add_run().add_break()
            p.add_run(line)


def generate_study_guide_docx():
    doc = docx_base('QRYVANTA Study Guide', 'SQL linting, ASTs, query performance and source rule reference | Edition 1.1')
    doc.add_heading('How to use this guide', level=1)
    doc.add_paragraph('Read the concept chapters first, then use the SQL001-SQL015 and P001-P022 chapters as source-specific reference sheets. Rule alerts are evidence to investigate, not automatic instructions. All practice questions in this bundle are original and are not represented as actual past-paper questions.')
    doc.add_heading('Contents', level=1)
    for title in [s['title'] for s in LEARN_SECTIONS] + ['SQL lint rules SQL001-SQL015', 'Performance rules P001-P022', 'Glossary', 'Credible references']:
        doc.add_paragraph(title, style='List Bullet')
    for section in LEARN_SECTIONS:
        doc.add_heading(section['title'], level=1)
        add_paragraphs(doc, section['body'])
    doc.add_heading('SQL lint rules SQL001-SQL015', level=1)
    doc.add_paragraph('These summaries follow the current implementation in linter/rules.py. A rule can be syntax-based, can have narrow trigger conditions, and can have false positives/negatives on unfamiliar dialect syntax.')
    for rule_id, title, severity, trigger, explanation, recommendation in SQL_RULES:
        doc.add_heading(f'{rule_id} - {title} [{severity}]', level=2)
        for label, value in [('Code trigger', trigger), ('Why it matters', explanation), ('Recommended response', recommendation)]:
            p = doc.add_paragraph()
            p.add_run(label + ': ').bold = True
            p.add_run(value)
    doc.add_heading('Performance rules P001-P022', level=1)
    doc.add_paragraph('These summaries follow analyzer/performance_rules.py and the normalized context/plan data it reads. Some checks require plan nodes or schema metadata. Source thresholds named in the implementation include LARGE_ROWS=10000, INDEX_FRACTION_THRESHOLD=0.30, ROW_EXPLOSION_FACTOR=2.0, LOW_CARDINALITY_MAX_DISTINCT=10, COVERING_MIN_ROWS=1000 and COVERING_MAX_COLUMNS=4.')
    for rule_id, title, severity, trigger, explanation, recommendation in PERFORMANCE_RULES:
        doc.add_heading(f'{rule_id} - {title} [{severity}]', level=2)
        for label, value in [('Code trigger', trigger), ('Why it matters', explanation), ('Recommended response', recommendation)]:
            p = doc.add_paragraph()
            p.add_run(label + ': ').bold = True
            p.add_run(value)
    doc.add_heading('Glossary', level=1)
    for key, value in sorted(GLOSSARY.items()):
        p = doc.add_paragraph()
        p.add_run(key + ': ').bold = True
        p.add_run(value)
    doc.add_heading('Credible references', level=1)
    for name, url, description in REFERENCES:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(name + ' - ').bold = True
        p.add_run(description + ' ' + url)
    path = OUTPUT / 'SQL_Linting_Performance_Study_Guide.docx'
    doc.save(path)
    return path


def make_pdf_styles():
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.units import mm
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='CoverTitleQ', parent=styles['Title'], fontName='Helvetica-Bold', fontSize=23, leading=29, alignment=TA_CENTER, textColor=colors.HexColor('#173B65'), spaceAfter=10))
    styles.add(ParagraphStyle(name='SubQ', parent=styles['Normal'], fontSize=10, leading=14, alignment=TA_CENTER, textColor=colors.HexColor('#526176'), spaceAfter=16))
    styles.add(ParagraphStyle(name='H1Q', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=15, leading=19, textColor=colors.HexColor('#204B78'), spaceBefore=12, spaceAfter=6, keepWithNext=True))
    styles.add(ParagraphStyle(name='H2Q', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=10.5, leading=13, textColor=colors.HexColor('#204B78'), spaceBefore=8, spaceAfter=4, keepWithNext=True))
    styles.add(ParagraphStyle(name='BodyQ', parent=styles['BodyText'], fontName='Helvetica', fontSize=8.6, leading=11.5, spaceAfter=4, wordWrap='CJK'))
    styles.add(ParagraphStyle(name='OptionQ', parent=styles['BodyText'], fontName='Helvetica', fontSize=8.4, leading=10.5, leftIndent=12, firstLineIndent=-9, spaceAfter=2, wordWrap='CJK'))
    styles.add(ParagraphStyle(name='AnswerQ', parent=styles['BodyText'], fontName='Helvetica-Bold', fontSize=8.5, leading=11, textColor=colors.HexColor('#12643F'), spaceBefore=3, spaceAfter=3, wordWrap='CJK'))
    styles.add(ParagraphStyle(name='ExplanationQ', parent=styles['BodyText'], fontName='Helvetica', fontSize=8.3, leading=10.8, leftIndent=8, borderPadding=5, backColor=colors.HexColor('#F1F5F9'), spaceAfter=7, wordWrap='CJK'))
    return styles


def add_pdf_header_footer(canvas, doc):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    canvas.saveState()
    width, height = A4
    canvas.setStrokeColor(colors.HexColor('#D8E1EB'))
    canvas.line(doc.leftMargin, 13*mm, width - doc.rightMargin, 13*mm)
    canvas.setFont('Helvetica', 7.5)
    canvas.setFillColor(colors.HexColor('#617086'))
    canvas.drawString(doc.leftMargin, 8*mm, 'QRYVANTA | Query Intelligence Studio')
    canvas.drawRightString(width - doc.rightMargin, 8*mm, f'Page {doc.page}')
    canvas.restoreState()



def generate_study_guide_pdf():
    """Create the matching PDF study guide without requiring LibreOffice."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, KeepTogether

    styles = make_pdf_styles()
    path = OUTPUT / 'SQL_Linting_Performance_Study_Guide.pdf'
    doc = SimpleDocTemplate(
        str(path), pagesize=A4, rightMargin=17*mm, leftMargin=17*mm,
        topMargin=17*mm, bottomMargin=19*mm,
        title='QRYVANTA SQL Linting and Performance Study Guide', author='QRYVANTA',
    )
    story = [
        Spacer(1, 32*mm),
        Paragraph('QRYVANTA Study Guide', styles['CoverTitleQ']),
        Paragraph('SQL linting, ASTs, query performance and source rule reference | Edition 1.1', styles['SubQ']),
        Paragraph('A practical offline course and source-aligned rule reference for Query Intelligence Studio. Rule alerts are evidence to investigate, not unconditional instructions to rewrite a query.', styles['BodyQ']),
        Spacer(1, 8*mm), Paragraph('How to use this guide', styles['H1Q']),
        Paragraph('Read the concept chapters first, then use the SQL001-SQL015 and P001-P022 sections as implementation-specific reference sheets. The rule summaries reflect the current source and its trigger conditions. Practice questions in this bundle are original, not claimed past-paper questions.', styles['BodyQ']),
        Paragraph('Contents', styles['H1Q']),
    ]
    contents = [section['title'] for section in LEARN_SECTIONS] + ['SQL lint rules SQL001-SQL015', 'Performance rules P001-P022', 'Glossary', 'Credible references']
    for title in contents:
        story.append(Paragraph('• ' + escape(title), styles['BodyQ']))
    story.append(PageBreak())

    for section in LEARN_SECTIONS:
        story.append(Paragraph(escape(section['title']), styles['H1Q']))
        for paragraph in section['body'].split('\n\n'):
            body = escape(paragraph.strip()).replace('\n', '<br/>')
            if body:
                story.append(Paragraph(body, styles['BodyQ']))

    story.append(Paragraph('SQL lint rules SQL001-SQL015', styles['H1Q']))
    story.append(Paragraph('These summaries follow the current implementation in linter/rules.py. A rule may use syntax-oriented heuristics and can have false positives or false negatives on uncommon dialect syntax.', styles['BodyQ']))
    for rule_id, title, severity, trigger, explanation, recommendation in SQL_RULES:
        elements = [Paragraph(escape(f'{rule_id} - {title} [{severity}]'), styles['H2Q'])]
        for label, value in [('Code trigger', trigger), ('Why it matters', explanation), ('Recommended response', recommendation)]:
            elements.append(Paragraph(f'<b>{escape(label)}:</b> {escape(value)}', styles['BodyQ']))
        story.append(KeepTogether(elements))

    story.append(Paragraph('Performance rules P001-P022', styles['H1Q']))
    story.append(Paragraph('These summaries follow analyzer/performance_rules.py and the normalized plan/context data it reads. Some checks require plan nodes or schema metadata. Source thresholds include LARGE_ROWS=10000, INDEX_FRACTION_THRESHOLD=0.30, ROW_EXPLOSION_FACTOR=2.0, LOW_CARDINALITY_MAX_DISTINCT=10, COVERING_MIN_ROWS=1000 and COVERING_MAX_COLUMNS=4.', styles['BodyQ']))
    for rule_id, title, severity, trigger, explanation, recommendation in PERFORMANCE_RULES:
        elements = [Paragraph(escape(f'{rule_id} - {title} [{severity}]'), styles['H2Q'])]
        for label, value in [('Code trigger', trigger), ('Why it matters', explanation), ('Recommended response', recommendation)]:
            elements.append(Paragraph(f'<b>{escape(label)}:</b> {escape(value)}', styles['BodyQ']))
        story.append(KeepTogether(elements))

    story.append(Paragraph('Glossary', styles['H1Q']))
    for key, value in sorted(GLOSSARY.items()):
        story.append(Paragraph(f'<b>{escape(key)}:</b> {escape(value)}', styles['BodyQ']))
    story.append(Paragraph('Credible references', styles['H1Q']))
    for name, url, description in REFERENCES:
        story.append(Paragraph(f'• <b>{escape(name)}</b> — {escape(description)} {escape(url)}', styles['BodyQ']))
    doc.build(story, onFirstPage=add_pdf_header_footer, onLaterPages=add_pdf_header_footer)
    return path


def generate_quiz_pdf():
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, KeepTogether
    styles = make_pdf_styles()
    path = OUTPUT / 'QRYVANTA_Quiz_Questions_Answers.pdf'
    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=17*mm, leftMargin=17*mm, topMargin=17*mm, bottomMargin=19*mm, title='QRYVANTA Quiz Questions, Answers and Explanations', author='QRYVANTA')
    story = [Spacer(1, 34*mm), Paragraph('QRYVANTA Quiz Bank', styles['CoverTitleQ']), Paragraph(f'{len(QUESTION_BANK)} original questions | Answers and explanations | Difficulty: Easy / Medium / Hard', styles['SubQ']),
             Paragraph('This document is generated from the same QUESTION_BANK as the in-app offline quiz. The document preserves canonical question/option order for study; the app randomizes question and option order in each attempt. These are original practice questions, not claimed past-paper questions.', styles['BodyQ']), Spacer(1, 7*mm), Paragraph('How to study', styles['H1Q']), Paragraph('Try each question before checking the answer. Use the explanation to revise the relevant idea or rule ID. Performance-rule questions describe the current QRYVANTA source implementation and the metadata it relies on.', styles['BodyQ']), PageBreak()]
    for number, q in enumerate(QUESTION_BANK, 1):
        group = [Paragraph(f"{number}. [{escape(q['id'])}] {escape(q['question'])}", styles['H2Q']), Paragraph(f"Topic: {escape(q['topic'])} | Difficulty: {escape(q['difficulty'])}", styles['BodyQ'])]
        for idx, option in enumerate(q['options']):
            group.append(Paragraph(f"{chr(65+idx)}. {escape(option)}", styles['OptionQ']))
        group.append(Paragraph(f"Correct answer: {escape(q['answer'])}", styles['AnswerQ']))
        group.append(Paragraph(f"Explanation: {escape(q['explanation'])}", styles['ExplanationQ']))
        story.append(KeepTogether(group))
    doc.build(story, onFirstPage=add_pdf_header_footer, onLaterPages=add_pdf_header_footer)
    return path


def generate_competitive_docx(solutions: bool):
    from docx.shared import Pt
    suffix = 'Solutions' if solutions else 'Questions'
    subtitle = (f'{len(QUESTION_BANK)} original solved practice questions | Answers and explanations included' if solutions else f'{len(QUESTION_BANK)} original SQL practice questions | Questions-only workbook')
    doc = docx_base(f'Competitive SQL Practice - {suffix}', subtitle)
    doc.add_paragraph('Attribution note: these are original QRYVANTA SQL practice questions written for revision. They are not presented as actual past-paper questions, and no exam/year label is asserted without a verified source.')
    for number, q in enumerate(QUESTION_BANK, 1):
        # Keep one compact question group together to avoid splitting an MCQ
        # across pages when Word/LibreOffice lays out the workbook.
        heading = doc.add_heading(f"{number}. [{q['id']}] {q['topic']} | {q['difficulty']}", level=2)
        group = [heading]
        group.append(doc.add_paragraph(q['question']))
        for letter, option in zip('ABCD', q['options']):
            group.append(doc.add_paragraph(f'{letter}. {option}', style='List Bullet'))
        if solutions:
            p = doc.add_paragraph()
            p.add_run('Correct answer: ').bold = True
            p.add_run(q['answer'])
            group.append(p)
            p = doc.add_paragraph()
            p.add_run('Explanation: ').bold = True
            p.add_run(q['explanation'])
            group.append(p)
        else:
            group.append(doc.add_paragraph('Answer: __________________________________________________________'))
        p = doc.add_paragraph()
        p.add_run('Topic tags: ').bold = True
        p.add_run(f"{q['topic']}; {q['difficulty']}; original practice")
        group.append(p)
        for paragraph in group[:-1]:
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = True
        group[-1].paragraph_format.keep_together = True
    path = OUTPUT / f'Competitive_SQL_Practice_{suffix}.docx'
    doc.save(path)
    return path


def main():
    paths = [generate_study_guide_docx(), generate_study_guide_pdf(), generate_quiz_pdf(), generate_competitive_docx(False), generate_competitive_docx(True)]
    for path in paths:
        print(f'{path.relative_to(ROOT)}\t{path.stat().st_size} bytes')

if __name__ == '__main__':
    main()
